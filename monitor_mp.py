# ==========================================================================
# MONITOR MERCADO PÚBLICO — Licitaciones + Compra Ágil + Órdenes de Compra
# Enfoque: materiales de construcción / ferretería, Región Metropolitana
# Versión GitHub Actions: lee el ticket desde la variable de entorno
# MP_TICKET, guarda todo en docs/ y genera docs/index.html (dashboard).
# ==========================================================================
import os
import requests
import pandas as pd
from datetime import datetime, timedelta, timezone
try:
    from zoneinfo import ZoneInfo
    TZ_CHILE = ZoneInfo("America/Santiago")
except Exception:
    TZ_CHILE = timezone(timedelta(hours=-3))
import time
import re
import unicodedata
 
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Patch
from openpyxl.drawing.image import Image as XLImage
 
# ----------------------- CONFIGURACIÓN GENERAL ----------------------------
TICKET = os.environ.get("MP_TICKET", "").strip()
if not TICKET:
    raise SystemExit(
        "Falta el ticket de Mercado Público.\n"
        "- En GitHub: Settings → Secrets and variables → Actions → "
        "New repository secret, nombre MP_TICKET.\n"
        "- En local/Colab: os.environ['MP_TICKET'] = 'TU_TICKET' antes de correr."
    )
 
BASE_V1 = "https://api.mercadopublico.cl/servicios/v1/publico"
URL_LICITACIONES = f"{BASE_V1}/licitaciones.json"
URL_OC = f"{BASE_V1}/ordenesdecompra.json"
BASE_V2 = "https://api2.mercadopublico.cl"
 
HEADERS_V1 = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"}
HEADERS_V2 = {"User-Agent": HEADERS_V1["User-Agent"], "ticket": TICKET}
 
# Activa/desactiva módulos
RUN_LICITACIONES = True
RUN_COMPRA_AGIL = True
RUN_ORDENES_COMPRA = True
 
# Órdenes de compra: cuántos días hacia atrás analizar (precios históricos)
OC_DIAS_ATRAS = 5
OC_ESTADO = "aceptada"   # aceptada = ya validadas por el proveedor
 
# Benchmark de precios: mínimo de líneas para publicar una referencia.
# Subirlo da menos referencias pero más sólidas.
MIN_LINEAS_BENCHMARK = 5
 
# Carpeta publicada por GitHub Pages
OUT_DIR = "docs"
os.makedirs(OUT_DIR, exist_ok=True)
EXCEL_PATH = os.path.join(OUT_DIR, "mercado_publico_construccion_RM.xlsx")
GRAFICO_PATH = os.path.join(OUT_DIR, "grafico_licitaciones.png")
HTML_PATH = os.path.join(OUT_DIR, "index.html")
HIST_PATH = os.path.join(OUT_DIR, "historico_oc_lineas.csv")
 
# ----------------------- KEYWORDS (motor común) ---------------------------
KEYWORDS_POR_CATEGORIA = {
    "cemento_hormigon": [
        "cemento", "cemento portland", "hormigon", "hormigon premezclado",
        "concreto", "mortero", "estuco", "aditivo hormigon",
        "hormigon armado", "pastina",
    ],
    "aridos": [
        "arido", "aridos", "gravilla", "grava", "arena", "arena gruesa",
        "arena fina", "ripio", "base estabilizada", "chancado",
        "polvo de piedra",
    ],
    "acero_metal": [
        "fierro", "acero", "fierro estriado", "barra de acero",
        "malla acma", "malla electrosoldada", "perfil metalico",
        "perfiles metalicos", "perfil estructural", "estructura metalica",
        "plancha de acero", "tubo metalico", "tubo galvanizado",
        "alambre negro", "alambre galvanizado",
    ],
    "madera": [
        "madera construccion", "madera de construccion", "madera aserrada",
        "pino oregon", "pino radiata", "madera impregnada",
        "tablero terciado", "contrachapado", "moldura de madera",
        "listones madera", "cerchas", "vigas de madera",
    ],
    "zinc_techumbre": [
        "plancha de zinc", "planchas de zinc", "zinc alum", "zincalum",
        "plancha ondulada", "cubierta metalica", "canaleta lluvia",
        "teja asfaltica", "membrana asfaltica", "impermeabilizante",
        "fieltro asfaltico",
    ],
    "albanileria": [
        "ladrillo", "ladrillo fiscal", "bloque de hormigon",
        "block hormigon", "yeso", "yeso carton", "volcanita",
        "panel yeso carton", "cal hidratada", "revestimiento muro",
    ],
    "pinturas_terminaciones": [
        "pintura construccion", "pintura latex", "pintura esmalte",
        "pintura anticorrosiva", "barniz", "sellador acrilico",
        "silicona sellante", "ceramica pavimento", "porcelanato",
        "piso flotante", "piso vinilico",
    ],
    "plomeria_sanitario": [
        "tuberia pvc", "caneria", "fittings pvc", "llave de paso",
        "artefacto sanitario",
    ],
    "electricidad": [
        "cable electrico", "canalizacion electrica", "tablero electrico",
        "conduit electrico",
    ],
    "ferreteria_general": [
        "ferreteria", "herramientas construccion",
        "material de construccion", "materiales de construccion",
        "insumos de construccion", "insumos construccion",
    ],
    "aislacion": [
        "lana de vidrio", "poliestireno expandido", "aislante termico",
        "aislacion acustica", "espuma poliuretano",
    ],
}
 
EXCLUDE_KEYWORDS = [
    "cemento oseo", "cemento dental", "arena sanitaria", "arena para gatos",
    "acero quirurgico", "malla quirurgica", "pintura facial", "pintura corporal",
    # Falsos positivos del rubro salud detectados en producción
    "aposito", "fluoruro", "monodosis", "curacion avanzada",
    "uso clinico", "uso dental",
]
 
CA_QUERIES = [
    "materiales de construccion", "ferreteria", "cemento", "fierro",
    "madera", "pintura", "hormigon", "ladrillo", "yeso", "zinc",
]
 
 
def normalizar(texto):
    # A prueba de None y NaN (los CSV releídos convierten vacíos en NaN float)
    if texto is None or (isinstance(texto, float) and pd.isna(texto)):
        return ""
    texto = str(texto).lower()
    texto = unicodedata.normalize("NFKD", texto)
    return "".join(c for c in texto if not unicodedata.combining(c))
 
 
KEYWORDS = sorted({
    normalizar(k)
    for lista in KEYWORDS_POR_CATEGORIA.values()
    for k in lista
})
EXCLUDE_KEYWORDS_NORM = [normalizar(k) for k in EXCLUDE_KEYWORDS]
 
 
def matchea_construccion(nombre):
    nombre_norm = normalizar(nombre)
    if any(ex in nombre_norm for ex in EXCLUDE_KEYWORDS_NORM):
        return False
    return any(re.search(rf"\b{re.escape(k)}\b", nombre_norm) for k in KEYWORDS)
 
 
def categorias_match(nombre):
    nombre_norm = normalizar(nombre)
    cats = []
    for cat, lista in KEYWORDS_POR_CATEGORIA.items():
        if any(re.search(rf"\b{re.escape(normalizar(k))}\b", nombre_norm) for k in lista):
            cats.append(cat)
    return cats
 
 
# ==========================================================================
# MOTOR DE PRECIOS — limpieza y normalización de líneas de órdenes de compra
# --------------------------------------------------------------------------
# Por qué existe: en Mercado Público el campo "Producto" es el nombre del
# catálogo ONU (UNSPSC) que elige el comprador, y suele no corresponder a lo
# que realmente se compró (bajo "Cemento" aparecen lijas, codos PVC, brochas).
# El producto real está escrito en EspecificacionComprador. Además la API NO
# devuelve unidad de medida (campo 100% vacío), así que el formato — saco de
# 25 kg, m3, plancha, galón — también se deduce de ese texto.
# ==========================================================================
 
# --- 1. Líneas "canasta": un carro de compras completo en una sola línea ---
# Ej: "MATERIALES DE FERRETERÍA SEGÚN LISTADO N°8 ADJUNTO A LA OC" por $588.104.
# No son un precio unitario y arruinan cualquier promedio.
PATRONES_CANASTA = [
    r"segun (el )?(listado|detalle|solicitud|requerimiento|memorandum|memo|"
    r"adjunto|ficha|anexo|orden|sp n|s\.p\.|pedido|cotizacion|presupuesto|"
    r"especificacion)",
    r"(listado|detalle|nomina|planilla|set|kit|lote|paquete) "
    r"(adjunto|anexo|n°|nro|numero)",
    r"adjunt[oa]s? a la oc", r"archivo adjunto", r"documento adjunto",
    r"varios (materiales|articulos|productos|items)",
    r"materiales (de |para )?(construccion|ferreteria|varios)",
    r"ayuda social", r"caso social", r"orasmi",
    r"segun (anexo|cuadro|tabla)",
]
RE_CANASTA = re.compile("|".join(PATRONES_CANASTA))
 
# Umbral sobre el cual una línea de cantidad 1 se considera suma alzada
UMBRAL_CANASTA_CLP = 150000
 
 
def es_linea_canasta(espec, cantidad, precio):
    """True si la línea agrupa varios productos en vez de ser un precio unitario."""
    e = normalizar(espec)
    if RE_CANASTA.search(e):
        return True
    try:
        if float(cantidad) == 1 and float(precio) > UMBRAL_CANASTA_CLP:
            return True
    except (TypeError, ValueError):
        pass
    # Varios saltos de línea = lista de productos escrita en un solo campo
    if str(espec).count("\n") >= 2:
        return True
    return False
 
 
# --- 2. Formato de venta, deducido del texto del comprador ---
# El orden importa: lo más específico primero (saco 25 kg antes que kg).
FORMATOS = [
    # En construcción el formato de 25 kg es el saco, se nombre o no la palabra
    ("saco 25 kg",    r"\b(saco|bolsa|bsa)s?\b.{0,25}\b25\s*(kg|kilo|k\b)|"
                      r"25\s*(kg|kilos?)\b.{0,20}\b(saco|bolsa)|"
                      r"\b25\s*(kg|kilos?)\b"),
    ("saco 42,5 kg",  r"\b(saco|bolsa)s?\b.{0,25}\b42[,.]?5\s*(kg|kilo)"),
    ("saco 20 kg",    r"\b(saco|bolsa)s?\b.{0,25}\b20\s*(kg|kilo)|"
                      r"20\s*kg\b.{0,20}\b(saco|bolsa)"),
    ("saco 1 kg",     r"\b(saco|bolsa)s?\b.{0,20}\b1\s*(kg|kilo)\b"),
    ("pallet",        r"\bpallet|\bpaleta\b"),
    ("galon",         r"\bgal[oó]n(es)?\b|\bgl\b"),
    ("litro",         r"\b\d+\s*(lt|lts|litros?)\b"),
    ("m3",            r"\bm\s?3\b|\bm³\b|metros?\s*c[uú]bicos?"),
    ("m2",            r"\bm\s?2\b|\bm²\b|metros?\s*cuadrados?"),
    ("metro lineal",  r"\bml\b|metros?\s*lineal|\bmt?s?\b(?!\s*\d)"),
    ("kg",            r"\b\d+\s*(kg|kilos?)\b"),
    ("plancha",       r"\bplancha(s)?\b"),
    ("rollo",         r"\brollo(s)?\b"),
    ("tira/barra",    r"\b(tira|barra|vara)s?\b"),
    ("caja",          r"\bcaja(s)?\b"),
    ("juego/set",     r"\b(juego|set|kit)\b"),
]
RE_FORMATOS = [(n, re.compile(p)) for n, p in FORMATOS]
 
 
def formato_venta(espec):
    """Saco de 25 kg, m3, plancha, galón... deducido del texto; 'unidad' si no hay pista."""
    e = normalizar(espec)
    for nombre, rx in RE_FORMATOS:
        if rx.search(e):
            return nombre
    return "unidad"
 
 
# --- 3. Material real, leído del texto (no de la clasificación ONU) ---
MATERIALES = {
    "cemento":         r"\bcemento\b",
    "hormigon":        r"\bhormig[oó]n\b|\bconcreto\b",
    "arena":           r"\barena\b",
    "gravilla/ripio":  r"\bgravilla\b|\bripio\b|\bgrava\b",
    "yeso/volcanita":  r"\byeso\b|\bvolcanita\b|yeso[- ]?carton",
    "ladrillo":        r"\bladrillo",
    "bloque":          r"\bbloque",
    "plancha zinc":    r"\bzinc\b|zincalum|\bcinc\b",
    "fierro/acero":    r"\bfierro\b|\bacero\b|barra.{0,10}(estriada|lisa)|malla acma",
    "madera pino":     r"\bpino\b|\bmadera\b|\btabla\b|\bterciado\b|\bosb\b|\bmdf\b",
    "pintura latex":   r"\bl[aá]tex\b",
    "pintura esmalte": r"\besmalte\b",
    "tubo pvc":        r"\bpvc\b|ca[ñn]eria",
    "clavos":          r"\bclavo",
    "tornillos":       r"\btornillo",
    "brocha/rodillo":  r"\bbrocha|\brodillo",
    "silicona/sello":  r"\bsilicona\b|\bsellante\b|\bsellador\b",
    "impermeabilizante": r"impermeabil|\bsika\b",
    "lija":            r"\blija",
}
RE_MATERIALES = [(n, re.compile(p)) for n, p in MATERIALES.items()]
 
 
def material_real(espec):
    """Material identificado en el texto del comprador, o None si no se reconoce."""
    e = normalizar(espec)
    for nombre, rx in RE_MATERIALES:
        if rx.search(e):
            return nombre
    return None
 
 
# --- 4. Dimensión (2x4, 18mm, 3/4...) — clave en madera, fierro y tubería ---
RE_DIMENSION = re.compile(
    r"(\d+\s?[x×]\s?\d+(\s?[x×]\s?\d+([,.]\d+)?)?)|(\b\d+\s?mm\b)|(\b\d+/\d+\b)"
)
MATERIALES_CON_DIMENSION = {
    "madera pino", "fierro/acero", "tubo pvc", "plancha zinc", "yeso/volcanita",
}
 
 
def dimension(espec, material):
    """Medida del producto cuando el material la necesita para ser comparable."""
    if material not in MATERIALES_CON_DIMENSION:
        return ""
    m = RE_DIMENSION.search(normalizar(espec))
    return m.group(0).replace(" ", "") if m else ""
 
 
def fecha_hora_legible(valor):
    """ISO-8601 de la API -> 'dd-mm-aaaa HH:MM' en hora de Chile. '' si no hay dato."""
    if not valor:
        return ""
    try:
        f = datetime.fromisoformat(str(valor).replace("Z", "+00:00"))
    except (ValueError, TypeError):
        return str(valor)
    if f.tzinfo is not None:
        f = f.astimezone(TZ_CHILE)
    return f.strftime("%d-%m-%Y %H:%M")
 
 
def recortar_outliers(serie):
    """Descarta valores fuera del rango intercuartil ampliado (regla 1,5 x IQR)."""
    if len(serie) < 5:
        return serie
    q1, q3 = serie.quantile(0.25), serie.quantile(0.75)
    iqr = q3 - q1
    return serie[(serie >= q1 - 1.5 * iqr) & (serie <= q3 + 1.5 * iqr)]
 
 
# ----------------------- HELPERS HTTP -------------------------------------
def get_v1(url, params, intentos=4, timeout=20):
    """GET a la API v1 (ticket como parámetro GET)."""
    params = dict(params, ticket=TICKET)
    for intento in range(intentos):
        try:
            r = requests.get(url, params=params, headers=HEADERS_V1, timeout=timeout)
            if r.status_code == 429:
                print("    (v1) Rate limit, esperando 15s...")
                time.sleep(15)
                continue
            r.raise_for_status()
            return r.json()
        except requests.exceptions.RequestException as e:
            print(f"    (v1) Intento {intento+1} falló: {e}")
            time.sleep(5)
    return None
 
 
def get_v2(path, params=None, intentos=3, timeout=40):
    """GET a la API v2 Compra Ágil (ticket en header). Maneja cuota diaria."""
    for intento in range(intentos):
        try:
            r = requests.get(f"{BASE_V2}{path}", params=params,
                             headers=HEADERS_V2, timeout=timeout)
            if r.status_code == 429:
                print("    (v2) CUOTA DIARIA AGOTADA — módulo Compra Ágil detenido.")
                return "QUOTA"
            if r.status_code in (401, 403):
                print("    (v2) Ticket no autorizado para API v2. "
                      "Solicita ticket en https://www.chilecompra.cl/api/")
                return "AUTH"
            r.raise_for_status()
            data = r.json()
            if data.get("success") == "OK":
                return data.get("payload")
            print(f"    (v2) Error API: {data.get('errors')}")
            return None
        except requests.exceptions.RequestException as e:
            print(f"    (v2) Intento {intento+1} falló: {e}")
            time.sleep(5)
    return None
 
 
# ==========================================================================
# MÓDULO 1: LICITACIONES
# ==========================================================================
df_lic = pd.DataFrame()
df_lineas = pd.DataFrame()
 
if RUN_LICITACIONES:
    print("=" * 60)
    print("MÓDULO 1: LICITACIONES ACTIVAS")
    print("=" * 60)
    data = get_v1(URL_LICITACIONES, {"estado": "activas"})
    activas = data.get("Listado", []) if data else []
    print(f"Total activas en Chile: {len(activas)}")
 
    candidatas = [l for l in activas if matchea_construccion(l.get("Nombre") or "")]
    print(f"Candidatas por palabra clave: {len(candidatas)}")
 
    resultados = []
    hoy = datetime.now()
    for idx, lic in enumerate(candidatas):
        codigo = lic.get("CodigoExterno")
        print(f"  [{idx+1}/{len(candidatas)}] {codigo}")
        data = get_v1(URL_LICITACIONES, {"codigo": codigo}, timeout=15)
        detalle = (data.get("Listado") or [None])[0] if data else None
        if detalle:
            comprador = detalle.get("Comprador", {}) or {}
            region = (comprador.get("RegionUnidad") or "").strip()
            estado = detalle.get("Estado", "")
            fecha_cierre_str = (detalle.get("Fechas", {}) or {}).get("FechaCierre")
            dias_restantes = None
            if fecha_cierre_str:
                try:
                    dias_restantes = (datetime.fromisoformat(fecha_cierre_str) - hoy).days
                except ValueError:
                    pass
 
            if "Metropolitana" in region and estado == "Publicada":
                nombre_lic = detalle.get("Nombre")
                base = {
                    "Codigo": codigo,
                    "NombreLicitacion": nombre_lic,
                    "TipoLicitacion": detalle.get("Tipo", ""),
                    "Organismo": comprador.get("NombreOrganismo", ""),
                    "Comuna": comprador.get("ComunaUnidad", ""),
                    "DiasParaCierre": dias_restantes,
                    "FechaCierre": fecha_cierre_str,
                    "MontoEstimado": detalle.get("MontoEstimado"),
                    "Moneda": detalle.get("Moneda", ""),
                    "Categorias": ", ".join(categorias_match(nombre_lic or "")),
                }
                items = (detalle.get("Items", {}) or {}).get("Listado", []) or []
                if items:
                    for n, item in enumerate(items, start=1):
                        fila = base.copy()
                        fila.update({
                            "Linea": n,
                            "ItemNombre": item.get("NombreProducto", ""),
                            "ItemDescripcion": item.get("Descripcion", ""),
                            "Cantidad": item.get("Cantidad"),
                            "UnidadMedida": item.get("UnidadMedida", ""),
                            "CodigoONU": item.get("CodigoProducto", ""),
                            "CategoriaONU": item.get("Categoria", ""),
                        })
                        resultados.append(fila)
                else:
                    fila = base.copy()
                    fila.update({"Linea": None, "ItemNombre": "", "ItemDescripcion": "",
                                 "Cantidad": None, "UnidadMedida": "",
                                 "CodigoONU": "", "CategoriaONU": ""})
                    resultados.append(fila)
        time.sleep(2)
 
    df_lineas = pd.DataFrame(resultados)
    if not df_lineas.empty:
        df_lineas = df_lineas.sort_values(["DiasParaCierre", "Codigo", "Linea"])
        df_lic = df_lineas.drop_duplicates("Codigo")[
            ["Codigo", "NombreLicitacion", "TipoLicitacion", "Organismo", "Comuna",
             "DiasParaCierre", "FechaCierre", "MontoEstimado", "Moneda", "Categorias"]
        ].reset_index(drop=True)
        monto_total = df_lic["MontoEstimado"].fillna(0).sum()
        print(f"\nLicitaciones RM Publicadas: {len(df_lic)}")
        print(f"Monto estimado total: ${monto_total:,.0f}".replace(",", "."))
 
 
# ==========================================================================
# MÓDULO 2: COMPRA ÁGIL (API v2, RM = región 13)
# ==========================================================================
df_ca = pd.DataFrame()
df_ca_productos = pd.DataFrame()
 
if RUN_COMPRA_AGIL:
    print("\n" + "=" * 60)
    print("MÓDULO 2: COMPRA ÁGIL — RM, estado publicada")
    print("=" * 60)
    ca_items = {}
    abortar = False
 
    for q in CA_QUERIES:
        if abortar:
            break
        print(f"  Buscando: '{q}'")
        params = {"q": q, "region": 13, "estado": "publicada",
                  "tamano_pagina": 50, "numero_pagina": 1}
        while True:
            payload = get_v2("/v2/compra-agil", params)
            if payload in ("QUOTA", "AUTH"):
                abortar = True
                break
            if not payload:
                break
            for item in payload.get("items", []):
                ca_items[item["codigo"]] = item
            pag = payload.get("paginacion", {})
            if pag.get("numero_pagina", 1) >= pag.get("total_paginas", 1):
                break
            params["numero_pagina"] += 1
            time.sleep(1)
        time.sleep(1)
 
    print(f"  Compras Ágiles únicas encontradas: {len(ca_items)}")
 
    filas_ca, filas_prod = [], []
    hoy_utc = datetime.now(timezone.utc)
    for idx, (codigo, item) in enumerate(ca_items.items()):
        if abortar:
            break
        nombre = item.get("nombre", "")
        if not matchea_construccion(nombre):
            continue
        print(f"  [{idx+1}/{len(ca_items)}] Detalle {codigo}")
        det = get_v2(f"/v2/compra-agil/{codigo}")
        if det in ("QUOTA", "AUTH"):
            abortar = True
            break
        if not det:
            continue
 
        fechas = det.get("fechas", {}) or {}
        fecha_cierre = fechas.get("fecha_cierre")
        dias = None
        if fecha_cierre:
            try:
                fc = datetime.fromisoformat(fecha_cierre.replace("Z", "+00:00"))
                if fc.tzinfo is None:
                    fc = fc.replace(tzinfo=timezone(timedelta(hours=-4)))
                dias = (fc - hoy_utc).days
            except (ValueError, TypeError):
                pass
 
        presupuesto = det.get("presupuesto", {}) or {}
        inst = det.get("institucion", {}) or {}
        conv = det.get("convocatoria", {}) or {}
 
        # Convocatoria: la Compra Ágil puede tener dos llamados para ofertar.
        # estado_convocatoria: 1 = primer llamado, 2 = segundo llamado.
        # Las fechas de cierre de cada llamado solo vienen en el DETALLE.
        etapa = conv.get("estado_convocatoria")
        cierre_1 = conv.get("fecha_cierre_primer_llamado")
        cierre_2 = conv.get("fecha_cierre_segundo_llamado")
 
        filas_ca.append({
            "Codigo": codigo,
            "Nombre": det.get("nombre", nombre),
            "Descripcion": (det.get("descripcion") or "")[:300],
            "Estado": (det.get("estado", {}) or {}).get("glosa", ""),
            "Llamado": conv.get("descripcion", ""),
            "EtapaLlamado": etapa,
            "CierrePrimerLlamado": fecha_hora_legible(cierre_1),
            # Si no hay segundo llamado la API devuelve null: lo decimos con
            # palabras para que una celda vacía no se lea como dato faltante.
            "CierreSegundoLlamado": (fecha_hora_legible(cierre_2)
                                     or "No tiene segundo llamado"),
            "Organismo": inst.get("organismo_comprador", ""),
            "Region": inst.get("nombre_region", ""),
            "FechaCierre": fecha_hora_legible(fecha_cierre),
            "DiasParaCierre": dias,
            "MontoDisponibleCLP": presupuesto.get("monto_disponible_clp"),
            "OfertasRecibidas": (det.get("resumen", {}) or {}).get("total_ofertas_recibidas"),
            "Categorias": ", ".join(categorias_match(det.get("nombre", nombre))),
        })
        for prod in det.get("productos_solicitados", []) or []:
            filas_prod.append({
                "CodigoCA": codigo,
                "Producto": prod.get("nombre", ""),
                "Descripcion": prod.get("descripcion") or "",
                "Cantidad": prod.get("cantidad"),
                "Unidad": prod.get("unidad_medida", ""),
                "CodigoProducto": prod.get("codigo_producto", ""),
            })
        time.sleep(1)
 
    df_ca = pd.DataFrame(filas_ca)
    df_ca_productos = pd.DataFrame(filas_prod)
    if not df_ca.empty:
        df_ca = df_ca.sort_values("DiasParaCierre")
        print(f"\nCompras Ágiles RM de construcción: {len(df_ca)}")
        print(f"Monto disponible total: "
              f"${df_ca['MontoDisponibleCLP'].fillna(0).sum():,.0f}".replace(",", "."))
 
 
# ==========================================================================
# MÓDULO 3: ÓRDENES DE COMPRA — análisis de precios por producto + unidad
# ==========================================================================
df_oc = pd.DataFrame()
df_oc_lineas = pd.DataFrame()
df_oc_limpias = pd.DataFrame()
df_precios = pd.DataFrame()
 
if RUN_ORDENES_COMPRA:
    print("\n" + "=" * 60)
    print(f"MÓDULO 3: ÓRDENES DE COMPRA (últimos {OC_DIAS_ATRAS} días, estado {OC_ESTADO})")
    print("=" * 60)
    candidatas_oc = []
    for d in range(1, OC_DIAS_ATRAS + 1):
        fecha = (datetime.now() - timedelta(days=d)).strftime("%d%m%Y")
        print(f"  Consultando OCs del {fecha}...")
        data = get_v1(URL_OC, {"fecha": fecha, "estado": OC_ESTADO})
        listado = data.get("Listado", []) if data else []
        matches = [oc for oc in listado if matchea_construccion(oc.get("Nombre") or "")]
        print(f"    Total día: {len(listado)} | Candidatas construcción: {len(matches)}")
        candidatas_oc.extend(matches)
        time.sleep(2)
 
    print(f"  Total OCs candidatas: {len(candidatas_oc)}")
 
    filas_oc, filas_oc_lineas = [], []
    for idx, oc in enumerate(candidatas_oc):
        codigo = oc.get("Codigo")
        print(f"  [{idx+1}/{len(candidatas_oc)}] Detalle OC {codigo}")
        data = get_v1(URL_OC, {"codigo": codigo}, timeout=15)
        det = (data.get("Listado") or [None])[0] if data else None
        if det:
            comprador = det.get("Comprador", {}) or {}
            proveedor = det.get("Proveedor", {}) or {}
            region = (comprador.get("RegionUnidad") or "").strip()
            nombre_oc = det.get("Nombre", "")
            base = {
                "CodigoOC": codigo,
                "NombreOC": nombre_oc,
                "Estado": det.get("Estado", ""),
                "FechaEnvio": (det.get("Fechas", {}) or {}).get("FechaEnvio", ""),
                "Organismo": comprador.get("NombreOrganismo", ""),
                "Region": region,
                "Proveedor": proveedor.get("Nombre", ""),
                "RutProveedor": proveedor.get("RutSucursal", ""),
                "Moneda": det.get("Moneda", ""),
                "TotalNeto": det.get("TotalNeto"),
                "Total": det.get("Total"),
                "Categorias": ", ".join(categorias_match(nombre_oc)),
            }
            filas_oc.append(base)
            items = (det.get("Items", {}) or {}).get("Listado", []) or []
            for item in items:
                filas_oc_lineas.append({
                    "CodigoOC": codigo,
                    "FechaEnvio": base["FechaEnvio"],
                    "Region": region,
                    "Proveedor": proveedor.get("Nombre", ""),
                    "Producto": item.get("Producto", ""),
                    "EspecificacionComprador": (item.get("EspecificacionComprador") or "")[:200],
                    "CategoriaONU": item.get("Categoria", ""),
                    "CodigoProductoONU": item.get("CodigoProducto", ""),
                    "Cantidad": item.get("Cantidad"),
                    # La API usa distintos nombres para la unidad según la versión
                    "Unidad": (item.get("Unidad") or item.get("UnidadMedida")
                               or item.get("unidad") or ""),
                    "PrecioNeto": item.get("PrecioNeto"),
                    "TotalLinea": item.get("Total"),
                    "Moneda": item.get("Moneda", det.get("Moneda", "")),
                })
        time.sleep(2)
 
    df_oc = pd.DataFrame(filas_oc)
    df_oc_lineas = pd.DataFrame(filas_oc_lineas)
 
    # ---------- Histórico acumulado de líneas de OC (crece en cada corrida) ----
    # Se carga siempre, incluso si la corrida no trajo líneas nuevas: así el
    # benchmark sigue publicándose un fin de semana o si la API falla.
    hist = pd.DataFrame()
    try:
        if os.path.exists(HIST_PATH):
            hist = pd.read_csv(HIST_PATH)
        if not df_oc_lineas.empty:
            hist = (pd.concat([hist, df_oc_lineas], ignore_index=True)
                    if not hist.empty else df_oc_lineas.copy())
            hist = hist.drop_duplicates(
                subset=["CodigoOC", "Producto", "Cantidad", "PrecioNeto"])
            hist.to_csv(HIST_PATH, index=False)
        print(f"  Histórico acumulado: {len(hist)} líneas en {HIST_PATH}")
    except Exception as e:
        print(f"  (No se pudo actualizar el histórico: {e})")
        hist = df_oc_lineas.copy()
 
    # ---------- Benchmark de precios: material x formato x dimensión ----------
    base_precios = hist if not hist.empty else df_oc_lineas
    if not base_precios.empty:
        b = base_precios.copy()
        for col in ["Moneda", "Producto", "EspecificacionComprador", "Region"]:
            if col in b.columns:
                b[col] = b[col].fillna("").astype(str)
        b = b[b["Moneda"].isin(["CLP", "", "nan", "None"])]
        b["PrecioNeto"] = pd.to_numeric(b["PrecioNeto"], errors="coerce")
        b["Cantidad"] = pd.to_numeric(b["Cantidad"], errors="coerce")
        b = b[b["PrecioNeto"].notna() & (b["PrecioNeto"] > 0)]
        n_bruto = len(b)
 
        # Paso 1: descartar líneas canasta (suma alzada, no precio unitario)
        b["EsCanasta"] = [
            es_linea_canasta(e, c, p)
            for e, c, p in zip(b["EspecificacionComprador"], b["Cantidad"], b["PrecioNeto"])
        ]
        n_canasta = int(b["EsCanasta"].sum())
        b = b[~b["EsCanasta"]]
 
        # Paso 2: leer material, formato y dimensión del texto del comprador
        b["Material"] = b["EspecificacionComprador"].map(material_real)
        b["Formato"] = b["EspecificacionComprador"].map(formato_venta)
        b["Dimension"] = [
            dimension(e, m) for e, m in zip(b["EspecificacionComprador"], b["Material"])
        ]
        df_oc_limpias = b[b["Material"].notna()].copy()
 
        # Paso 3: agrupar y recortar outliers dentro de cada grupo
        filas_bench = []
        for (mat, fmt, dim), sub in df_oc_limpias.groupby(
                ["Material", "Formato", "Dimension"], dropna=False):
            if len(sub) < MIN_LINEAS_BENCHMARK:
                continue
            precios = recortar_outliers(sub["PrecioNeto"])
            if len(precios) < 3:
                continue
            prom = precios.mean()
            cv = (precios.std() / prom) if prom and len(precios) > 1 else 0
            filas_bench.append({
                "Material": mat,
                "Formato": fmt,
                "Dimension": dim,
                "NumOC": sub["CodigoOC"].nunique(),
                "NumLineas": len(precios),
                "Descartadas": len(sub) - len(precios),
                "PrecioP25": round(precios.quantile(0.25)),
                "PrecioMediana": round(precios.median()),
                # La API entrega precios netos; este es el equivalente bruto
                # para comparar contra precios de retail (que incluyen IVA)
                "MedianaConIVA": round(precios.median() * 1.19),
                "PrecioP75": round(precios.quantile(0.75)),
                "PrecioMin": round(precios.min()),
                "PrecioMax": round(precios.max()),
                "CV": round(cv, 2),
                "EjemploEspecificacion": str(
                    sub["EspecificacionComprador"].iloc[0])[:90],
            })
 
        df_precios = pd.DataFrame(filas_bench)
        if not df_precios.empty:
            df_precios["Confiabilidad"] = pd.cut(
                df_precios["CV"],
                bins=[-0.01, 0.3, 0.8, float("inf")],
                labels=["Alta", "Media", "Baja (revisar)"],
            )
            df_precios = df_precios.sort_values(
                ["Confiabilidad", "NumLineas"], ascending=[True, False]
            ).reset_index(drop=True)
 
        print(f"\nOCs de construcción encontradas: {len(df_oc)}")
        print(f"Líneas con precio (CLP, histórico): {n_bruto}")
        print(f"  - descartadas por ser canasta/suma alzada: {n_canasta}")
        print(f"  - con material identificado en la especificación: {len(df_oc_limpias)}")
        print(f"Referencias de precio (material x formato x dimensión): {len(df_precios)}")
        if not df_precios.empty:
            n_alta = int((df_precios["Confiabilidad"] == "Alta").sum())
            print(f"  de ellas, confiabilidad Alta: {n_alta}")
 
 
# ==========================================================================
# GRÁFICO (licitaciones) + EXPORTACIÓN A EXCEL
# ==========================================================================
hojas = {}
if not df_lic.empty:
    hojas["Licitaciones"] = df_lic
    hojas["Lic_Lineas"] = df_lineas
if not df_ca.empty:
    hojas["CompraAgil"] = df_ca
    if not df_ca_productos.empty:
        hojas["CA_Productos"] = df_ca_productos
if not df_oc.empty:
    hojas["OC_Resumen"] = df_oc
    hojas["OC_Lineas"] = df_oc_lineas
# Las hojas de precios salen del histórico, no de las OCs de hoy
if not df_precios.empty:
    hojas["Precios_Productos"] = df_precios
if not df_oc_limpias.empty:
    # Trazabilidad: las líneas que alimentan cada referencia de precio
    hojas["Precios_Detalle"] = df_oc_limpias[[
        "CodigoOC", "FechaEnvio", "Region", "Proveedor", "Material",
        "Formato", "Dimension", "Cantidad", "PrecioNeto",
        "EspecificacionComprador", "Producto",
    ]]
 
grafico_ok = False
if not df_lic.empty:
    df_plot = df_lic.copy()
    df_plot["MontoEstimado"] = df_plot["MontoEstimado"].fillna(0)
    df_plot["CategoriaPrincipal"] = (
        df_plot["Categorias"].fillna("sin_categoria")
        .apply(lambda c: c.split(",")[0].strip() if c else "sin_categoria")
    )
    if len(df_plot) > 30:
        df_plot = df_plot.nlargest(30, "MontoEstimado")
    df_plot = df_plot.sort_values("MontoEstimado")
 
    cats = df_plot["CategoriaPrincipal"].unique()
    cmap = plt.get_cmap("tab20", max(len(cats), 1))
    color_por_cat = {c: cmap(i) for i, c in enumerate(cats)}
 
    alto = max(4, 0.45 * len(df_plot))
    fig, ax = plt.subplots(figsize=(12, alto))
    barras = ax.barh(df_plot["Codigo"], df_plot["MontoEstimado"],
                     color=df_plot["CategoriaPrincipal"].map(color_por_cat))
    max_monto = df_plot["MontoEstimado"].max() or 1
    for barra, dias in zip(barras, df_plot["DiasParaCierre"]):
        et = f"{int(dias)}d" if pd.notna(dias) else "s/f"
        ax.text(barra.get_width() + max_monto * 0.01,
                barra.get_y() + barra.get_height() / 2, et, va="center", fontsize=8)
    ax.legend(handles=[Patch(color=color_por_cat[c], label=c) for c in cats],
              title="Categoría", fontsize=8, title_fontsize=9, loc="lower right")
    ax.set_xlabel("Monto estimado (CLP)")
    ax.set_ylabel("ID Licitación")
    ax.set_title("Licitaciones construcción RM: monto y días para cierre")
    ax.xaxis.set_major_formatter(
        plt.FuncFormatter(lambda x, _: f"${x:,.0f}".replace(",", ".")))
    plt.tight_layout()
    fig.savefig(GRAFICO_PATH, dpi=150)
    plt.close(fig)
    grafico_ok = True
 
if hojas:
    with pd.ExcelWriter(EXCEL_PATH, engine="openpyxl") as writer:
        for nombre, df_hoja in hojas.items():
            df_hoja.to_excel(writer, sheet_name=nombre, index=False)
        if "Licitaciones" in hojas:
            ws = writer.sheets["Licitaciones"]
            fila_total = len(df_lic) + 2
            ws.cell(row=fila_total, column=1, value="TOTAL")
            ws.cell(row=fila_total, column=8,
                    value=float(df_lic["MontoEstimado"].fillna(0).sum()))
        if grafico_ok:
            ws_g = writer.book.create_sheet("Grafico")
            ws_g.add_image(XLImage(GRAFICO_PATH), "B2")
    print(f"\n✅ Exportado a {EXCEL_PATH}")
 
 
# ==========================================================================
# DASHBOARD HTML (docs/index.html, publicado por GitHub Pages)
# ==========================================================================
def fmt_clp(v):
    try:
        if pd.isna(v):
            return ""
        return f"${v:,.0f}".replace(",", ".")
    except Exception:
        return v
 
 
def tabla_html(df, cols_monto=(), max_rows=None):
    if df.empty:
        return "<p class='vacio'>Sin resultados en esta corrida.</p>"
    d = df.copy() if max_rows is None else df.head(max_rows).copy()
    for c in cols_monto:
        if c in d.columns:
            d[c] = d[c].apply(fmt_clp)
    return d.to_html(index=False, border=0, classes="tabla", na_rep="")
 
 
ahora_cl = datetime.now(TZ_CHILE).strftime("%d-%m-%Y %H:%M")
n_lic = len(df_lic)
monto_lic = fmt_clp(df_lic["MontoEstimado"].fillna(0).sum()) if not df_lic.empty else "$0"
n_ca = len(df_ca)
monto_ca = fmt_clp(df_ca["MontoDisponibleCLP"].fillna(0).sum()) if not df_ca.empty else "$0"
n_oc = len(df_oc)
n_prod = len(df_precios)
n_alta_conf = (int((df_precios["Confiabilidad"] == "Alta").sum())
               if not df_precios.empty else 0)
 
html = f"""<!doctype html>
<html lang="es">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Monitor Mercado Público — Construcción RM</title>
<style>
  :root {{ --azul:#1a3c6e; --gris:#f4f6f8; --borde:#dde3ea; }}
  body {{ font-family: -apple-system, Segoe UI, Roboto, sans-serif; margin:0;
         background:var(--gris); color:#222; }}
  header {{ background:var(--azul); color:#fff; padding:24px 16px; }}
  header h1 {{ margin:0 0 4px; font-size:1.4rem; }}
  header p {{ margin:0; opacity:.85; font-size:.9rem; }}
  main {{ max-width:1200px; margin:0 auto; padding:16px; }}
  .kpis {{ display:flex; flex-wrap:wrap; gap:12px; margin:16px 0; }}
  .kpi {{ background:#fff; border:1px solid var(--borde); border-radius:10px;
          padding:14px 18px; flex:1; min-width:180px; }}
  .kpi .v {{ font-size:1.4rem; font-weight:700; color:var(--azul); }}
  .kpi .l {{ font-size:.8rem; color:#666; }}
  section {{ background:#fff; border:1px solid var(--borde); border-radius:10px;
             padding:16px; margin-bottom:20px; overflow-x:auto; }}
  h2 {{ font-size:1.1rem; color:var(--azul); margin-top:0; }}
  .tabla {{ border-collapse:collapse; width:100%; font-size:.82rem; }}
  .tabla th {{ background:var(--azul); color:#fff; padding:6px 8px;
               text-align:left; position:sticky; top:0; }}
  .tabla td {{ padding:5px 8px; border-bottom:1px solid var(--borde); }}
  .tabla tr:nth-child(even) {{ background:#fafbfc; }}
  img {{ max-width:100%; height:auto; }}
  .acciones a {{ display:inline-block; background:var(--azul); color:#fff;
                 text-decoration:none; padding:10px 18px; border-radius:8px;
                 margin-right:10px; font-size:.9rem; }}
  .vacio {{ color:#888; }}
  .nota {{ font-size:.82rem; color:#555; margin:0 0 12px; line-height:1.45; }}
  .leyenda {{ font-size:.82rem; color:#444; line-height:1.5; margin-top:14px;
              padding:12px 14px; background:#f7f9fb;
              border-left:3px solid var(--azul); border-radius:0 8px 8px 0; }}
  .leyenda ul {{ margin:8px 0 8px; padding-left:20px; }}
  .leyenda li {{ margin-bottom:4px; }}
  footer {{ text-align:center; color:#888; font-size:.8rem; padding:20px; }}
</style>
</head>
<body>
<header>
  <h1>Monitor Mercado Público — Materiales de Construcción (RM)</h1>
  <p>Última actualización: {ahora_cl} (hora de Chile)</p>
</header>
<main>
  <div class="kpis">
    <div class="kpi"><div class="v">{n_lic}</div><div class="l">Licitaciones RM publicadas</div></div>
    <div class="kpi"><div class="v">{monto_lic}</div><div class="l">Monto estimado licitaciones</div></div>
    <div class="kpi"><div class="v">{n_ca}</div><div class="l">Compras Ágiles RM publicadas</div></div>
    <div class="kpi"><div class="v">{monto_ca}</div><div class="l">Monto disponible Compra Ágil</div></div>
    <div class="kpi"><div class="v">{n_oc}</div><div class="l">OCs analizadas (precios)</div></div>
    <div class="kpi"><div class="v">{n_prod}</div><div class="l">Precios de referencia</div></div>
    <div class="kpi"><div class="v">{n_alta_conf}</div><div class="l">De confiabilidad alta</div></div>
  </div>
 
  <section class="acciones">
    <a href="mercado_publico_construccion_RM.xlsx">⬇ Descargar Excel completo</a>
    <a href="historico_oc_lineas.csv">⬇ Histórico de precios (CSV)</a>
  </section>
 
  <section>
    <h2>Licitaciones RM (estado Publicada)</h2>
    {tabla_html(df_lic, cols_monto=["MontoEstimado"])}
  </section>
 
  <section>
    <h2>Gráfico: monto y días para cierre</h2>
    {('<img src="grafico_licitaciones.png" alt="Gráfico licitaciones">' if grafico_ok else "<p class='vacio'>Sin gráfico en esta corrida.</p>")}
  </section>
 
  <section>
    <h2>Compras Ágiles RM (publicadas)</h2>
    {tabla_html(
        df_ca[["Codigo", "Nombre", "Organismo", "Llamado",
               "CierrePrimerLlamado", "CierreSegundoLlamado", "DiasParaCierre",
               "MontoDisponibleCLP", "OfertasRecibidas", "Categorias"]]
        if not df_ca.empty else df_ca,
        cols_monto=["MontoDisponibleCLP"])}
    <div class="leyenda">
      <b>Cómo leer los llamados.</b> Una Compra Ágil se puede ofertar en hasta
      dos llamados. Lo primero que hay que mirar es la columna
      <b>CierreSegundoLlamado</b>:
      <ul>
        <li><b>"No tiene segundo llamado"</b> → la única oportunidad de ofertar
        es el primer llamado. Si esa fecha ya pasó, quedó fuera de plazo.</li>
        <li><b>Una fecha y hora</b> → sí hay segundo llamado, y ese es tu plazo
        real para ofertar, aunque el primero ya esté cerrado.</li>
      </ul>
      La columna <b>Llamado</b> dice en cuál está abierta en este momento
      ("Primer llamado" o "Segundo llamado"), y <b>DiasParaCierre</b> cuenta los
      días hasta el cierre de ese llamado vigente — no hasta el primero.
      Las horas están en hora de Chile.
    </div>
  </section>
 
  <section>
    <h2>Precios de referencia (histórico acumulado)</h2>
    <p class="nota">Precio unitario en CLP por material y formato de venta.
    Usa la <b>mediana</b> como referencia. Los precios de Mercado Público son
    <b>netos</b>; la columna <b>MedianaConIVA</b> es el equivalente bruto, para
    comparar contra precios de retail. Se excluyen las líneas de suma alzada
    ("según listado adjunto") y los valores atípicos de cada grupo.
    Confiabilidad <b>Alta</b> = precios muy consistentes entre sí.</p>
    {tabla_html(
        df_precios[["Material", "Formato", "Dimension", "Confiabilidad",
                    "NumOC", "NumLineas", "PrecioP25", "PrecioMediana",
                    "MedianaConIVA", "PrecioP75", "CV", "EjemploEspecificacion"]]
        if not df_precios.empty else df_precios,
        cols_monto=["PrecioP25", "PrecioMediana", "MedianaConIVA",
                    "PrecioP75"], max_rows=60)}
  </section>
</main>
<footer>Fuente: API Mercado Público (ChileCompra) · Generado automáticamente con GitHub Actions</footer>
</body>
</html>"""
 
with open(HTML_PATH, "w", encoding="utf-8") as f:
    f.write(html)
print(f"✅ Dashboard generado en {HTML_PATH}")
 
